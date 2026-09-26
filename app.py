from flask import (
    Flask,
    render_template,
    request,
    jsonify,
    redirect,
    url_for,
    send_file
)

from flask_login import (
    LoginManager,
    UserMixin,
    login_user,
    login_required,
    logout_user,
    current_user
)

from werkzeug.security import (
    generate_password_hash,
    check_password_hash
)

import sqlite3
import json
import io

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment
)
from openpyxl.utils import get_column_letter


# =========================================================
# إعداد التطبيق
# =========================================================

app = Flask(__name__)

app.secret_key = "CHANGE-THIS-TO-A-LONG-RANDOM-SECRET-KEY"

login_manager = LoginManager()

login_manager.init_app(app)

login_manager.login_view = "login"


# =========================================================
# قاعدة البيانات
# =========================================================

DATABASE = "reports.db"


# =========================================================
# المستخدم
# =========================================================

class User(UserMixin):

    def __init__(
        self,
        user_id,
        username,
        password_hash
    ):

        self.id = user_id

        self.username = username

        self.password_hash = password_hash


# =========================================================
# تحميل المستخدم من قاعدة البيانات
# =========================================================

@login_manager.user_loader
def load_user(user_id):

    connection = sqlite3.connect(DATABASE)

    cursor = connection.cursor()

    cursor.execute("""
        SELECT
            id,
            username,
            password_hash

        FROM users

        WHERE id = ?
    """, (user_id,))

    row = cursor.fetchone()

    connection.close()

    if row is None:

        return None

    return User(
        row[0],
        row[1],
        row[2]
    )


# =========================================================
# إنشاء قاعدة البيانات
# =========================================================

def init_database():

    connection = sqlite3.connect(DATABASE)

    cursor = connection.cursor()


    # -----------------------------------------------------
    # جدول المستخدمين
    # -----------------------------------------------------

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS users (

            id INTEGER PRIMARY KEY AUTOINCREMENT,

            username TEXT UNIQUE NOT NULL,

            password_hash TEXT NOT NULL

        )
    """)


    # -----------------------------------------------------
    # جدول التقارير
    # -----------------------------------------------------

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS reports (

            report_date TEXT PRIMARY KEY,

            report_data TEXT NOT NULL

        )
    """)


    # -----------------------------------------------------
    # إنشاء المستخدم الأساسي
    # -----------------------------------------------------

    cursor.execute("""
        SELECT id

        FROM users

        WHERE username = ?
    """, ("admin",))

    existing_user = cursor.fetchone()


    if existing_user is None:

        password_hash = generate_password_hash(
            "123456"
        )

        cursor.execute("""
            INSERT INTO users
            (
                username,
                password_hash
            )

            VALUES (?, ?)
        """, (
            "admin",
            password_hash
        ))


    connection.commit()

    connection.close()


# =========================================================
# تسجيل الدخول
# =========================================================

@app.route(
    "/login",
    methods=["GET", "POST"]
)
def login():

    # إذا كان المستخدم مسجل دخول بالفعل
    if current_user.is_authenticated:

        return redirect(
            url_for("index")
        )


    error = None


    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()


        password = request.form.get(
            "password",
            ""
        )


        # -------------------------------------------------
        # البحث عن المستخدم
        # -------------------------------------------------

        connection = sqlite3.connect(
            DATABASE
        )

        cursor = connection.cursor()

        cursor.execute("""
            SELECT
                id,
                username,
                password_hash

            FROM users

            WHERE username = ?
        """, (username,))

        row = cursor.fetchone()

        connection.close()


        # -------------------------------------------------
        # المستخدم غير موجود
        # -------------------------------------------------

        if row is None:

            error = (
                "اسم المستخدم أو كلمة المرور غير صحيحة."
            )


        else:

            user_id = row[0]

            username_from_db = row[1]

            password_hash = row[2]


            # ---------------------------------------------
            # التحقق من كلمة المرور
            # ---------------------------------------------

            password_correct = check_password_hash(
                password_hash,
                password
            )


            if password_correct:

                user = User(
                    user_id,
                    username_from_db,
                    password_hash
                )


                # -----------------------------------------
                # إنشاء جلسة تسجيل الدخول
                # -----------------------------------------

                login_user(user)


                return redirect(
                    url_for("index")
                )


            else:

                error = (
                    "اسم المستخدم أو كلمة المرور غير صحيحة."
                )


    return render_template(
        "login.html",
        error=error
    )


# =========================================================
# تسجيل الخروج
# =========================================================

@app.route("/logout")
@login_required
def logout():

    logout_user()

    return redirect(
        url_for("login")
    )


# =========================================================
# الصفحة الرئيسية
# =========================================================

@app.route("/")
@login_required
def index():

    return render_template(
        "index.html"
    )


# =========================================================
# حفظ التقرير
# =========================================================

@app.route(
    "/api/save",
    methods=["POST"]
)
@login_required
def save_report():

    try:

        data = request.get_json()


        if not data:

            return jsonify({

                "success": False,

                "message":
                    "لم يتم إرسال بيانات"

            }), 400


        report_date = data.get(
            "date"
        )


        report_data = data.get(
            "data"
        )


        if not report_date:

            return jsonify({

                "success": False,

                "message":
                    "التاريخ مطلوب"

            }), 400


        if report_data is None:

            return jsonify({

                "success": False,

                "message":
                    "بيانات التقرير مطلوبة"

            }), 400


        connection = sqlite3.connect(
            DATABASE
        )

        cursor = connection.cursor()


        cursor.execute("""
            INSERT OR REPLACE INTO reports
            (
                report_date,
                report_data
            )

            VALUES (?, ?)
        """, (

            report_date,

            json.dumps(
                report_data,
                ensure_ascii=False
            )

        ))


        connection.commit()

        connection.close()


        return jsonify({

            "success": True,

            "message":
                "تم حفظ التقرير بنجاح"

        })


    except Exception as error:

        print(
            "SAVE ERROR:",
            error
        )


        return jsonify({

            "success": False,

            "message":
                str(error)

        }), 500


# =========================================================
# تحميل تقرير بتاريخ معين
# =========================================================

@app.route(
    "/api/load/<report_date>"
)
@login_required
def load_report(report_date):

    try:

        connection = sqlite3.connect(
            DATABASE
        )

        cursor = connection.cursor()


        cursor.execute("""
            SELECT report_data

            FROM reports

            WHERE report_date = ?
        """, (report_date,))


        result = cursor.fetchone()

        connection.close()


        if result is None:

            return jsonify({

                "success": False,

                "message":
                    "لا يوجد تقرير لهذا التاريخ"

            }), 404


        report_data = json.loads(
            result[0]
        )


        return jsonify({

            "success": True,

            "data": report_data

        })


    except Exception as error:

        print(
            "LOAD ERROR:",
            error
        )


        return jsonify({

            "success": False,

            "message":
                str(error)

        }), 500


# =========================================================
# قائمة التقارير السابقة
# =========================================================

@app.route(
    "/api/reports"
)
@login_required
def get_reports():

    try:

        connection = sqlite3.connect(
            DATABASE
        )

        cursor = connection.cursor()


        cursor.execute("""
            SELECT report_date

            FROM reports

            ORDER BY report_date DESC
        """)


        rows = cursor.fetchall()

        connection.close()


        dates = [

            row[0]

            for row in rows

        ]


        return jsonify({

            "success": True,

            "reports": dates

        })


    except Exception as error:

        print(
            "REPORTS ERROR:",
            error
        )


        return jsonify({

            "success": False,

            "message":
                str(error)

        }), 500


# =========================================================
# تصدير Excel
# =========================================================

@app.route(
    "/api/export-excel/<report_date>"
)
@login_required
def export_excel(report_date):

    try:

        # -------------------------------------------------
        # جلب التقرير
        # -------------------------------------------------

        connection = sqlite3.connect(
            DATABASE
        )

        cursor = connection.cursor()


        cursor.execute("""
            SELECT report_data

            FROM reports

            WHERE report_date = ?
        """, (report_date,))


        result = cursor.fetchone()

        connection.close()


        if result is None:

            return jsonify({

                "success": False,

                "message":
                    "لا يوجد تقرير محفوظ لهذا التاريخ"

            }), 404


        saved_data = json.loads(
            result[0]
        )


        # -------------------------------------------------
        # تعريف الأقسام
        # -------------------------------------------------

        reports = [

            {
                "name": "جريت ستيك",

                "columns": [
                    "طلبات توصيل طلبات",
                    "سنونو",
                    "كيتا - توصيل كيتا",
                    "ديليفرو توصيل المطعم",
                    "الكول سنتر"
                ],

                "branches": [
                    "حولي",
                    "مارينا مول",
                    "مشرف",
                    "صباح السالم",
                    "المنقف",
                    "سعد العبدالله",
                    "العارضية"
                ]
            },


            {
                "name": "ستيك اند فرايز",

                "columns": [
                    "طلبات توصيل طلبات",
                    "جاهز توصيل جاهز",
                    "كيتا - توصيل كيتا",
                    "ديليفرو توصيل المطعم",
                    "الكول سنتر"
                ],

                "branches": [
                    "حولي",
                    "صباح السالم",
                    "المنقف",
                    "سعد العبدالله",
                    "العارضية"
                ]
            },


            {
                "name": "برجر 8",

                "columns": [
                    "طلبات توصيل طلبات",
                    "جاهز توصيل جاهز",
                    "كيتا - توصيل كيتا",
                    "ديليفرو توصيل المطعم",
                    "الكول سنتر"
                ],

                "branches": [
                    "حولي",
                    "مارينا مول",
                    "مشرف",
                    "صباح السالم",
                    "المنقف",
                    "سعد العبدالله",
                    "العارضية"
                ]
            },


            {
                "name": "بيتزا روستيكا",

                "columns": [
                    "طلبات توصيل المطعم",
                    "جاهز توصيل جاهز",
                    "كيتا - توصيل كيتا",
                    "ديليفرو توصيل المطعم",
                    "الكول سنتر"
                ],

                "branches": [
                    "صباح السالم"
                ]
            }

        ]


        # -------------------------------------------------
        # إنشاء ملف Excel
        # -------------------------------------------------

        workbook = Workbook()

        worksheet = workbook.active

        worksheet.title = "تقرير المبيعات"


        # -------------------------------------------------
        # الألوان
        # -------------------------------------------------

        dark_blue = "111A2E"

        orange = "F5C89E"

        yellow = "FFEDC2"

        white = "FFFFFF"

        black = "000000"


        thin_side = Side(
            style="thin",
            color=black
        )


        border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=thin_side
        )


        # -------------------------------------------------
        # عنوان التقرير
        # -------------------------------------------------

        worksheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=11
        )


        title_cell = worksheet.cell(
            row=1,
            column=1
        )


        title_cell.value = (
            "تقرير المبيعات والطلبات - "
            + report_date
        )


        title_cell.font = Font(
            bold=True,
            size=18,
            color=white
        )


        title_cell.fill = PatternFill(
            "solid",
            fgColor=dark_blue
        )


        title_cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


        worksheet.row_dimensions[1].height = 35


        current_row = 3


        # -------------------------------------------------
        # الأقسام
        # -------------------------------------------------

        for report_index, report in enumerate(
            reports
        ):


            # ---------------------------------------------
            # اسم القسم
            # ---------------------------------------------

            worksheet.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row,
                end_column=11
            )


            restaurant_cell = worksheet.cell(
                row=current_row,
                column=1
            )


            restaurant_cell.value = report["name"]


            restaurant_cell.font = Font(
                bold=True,
                size=15,
                color=white
            )


            restaurant_cell.fill = PatternFill(
                "solid",
                fgColor=dark_blue
            )


            restaurant_cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )


            current_row += 1


            # ---------------------------------------------
            # رأس الجدول
            # ---------------------------------------------

            worksheet.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row + 1,
                end_column=1
            )


            branch_header = worksheet.cell(
                row=current_row,
                column=1
            )


            branch_header.value = "الفرع"


            excel_column = 2


            for company in report["columns"]:


                worksheet.merge_cells(
                    start_row=current_row,
                    start_column=excel_column,
                    end_row=current_row,
                    end_column=excel_column + 1
                )


                cell = worksheet.cell(
                    row=current_row,
                    column=excel_column
                )


                cell.value = company


                worksheet.cell(
                    row=current_row + 1,
                    column=excel_column
                ).value = "الطلبات"


                worksheet.cell(
                    row=current_row + 1,
                    column=excel_column + 1
                ).value = "المبلغ"


                excel_column += 2


            # تنسيق الرأس

            for row in range(
                current_row,
                current_row + 2
            ):

                for col in range(1, 12):

                    cell = worksheet.cell(
                        row=row,
                        column=col
                    )


                    cell.border = border


                    cell.fill = PatternFill(
                        "solid",
                        fgColor=orange
                    )


                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical="center",
                        wrap_text=True
                    )


                    cell.font = Font(
                        bold=True
                    )


            current_row += 2


            # ---------------------------------------------
            # بيانات الفروع
            # ---------------------------------------------

            report_data = saved_data[
                report_index
            ]


            for row_index, branch in enumerate(
                report["branches"]
            ):

                branch_cell = worksheet.cell(
                    row=current_row,
                    column=1
                )


                branch_cell.value = branch


                branch_cell.font = Font(
                    bold=True
                )


                branch_cell.fill = PatternFill(
                    "solid",
                    fgColor="F4F5F7"
                )


                branch_cell.border = border


                branch_cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )


                row_data = report_data[
                    row_index
                ]


                for data_index, value in enumerate(
                    row_data
                ):

                    cell = worksheet.cell(
                        row=current_row,
                        column=data_index + 2
                    )


                    cell.value = value


                    cell.border = border


                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical="center"
                    )


                    if data_index % 2 == 1:

                        cell.number_format = "0.000"


                current_row += 1


            # ---------------------------------------------
            # الإجمالي
            # ---------------------------------------------

            total_cell = worksheet.cell(
                row=current_row,
                column=1
            )


            total_cell.value = "الإجمالي"


            total_cell.font = Font(
                bold=True
            )


            total_cell.fill = PatternFill(
                "solid",
                fgColor=yellow
            )


            total_cell.border = border


            total_cell.alignment = Alignment(
                horizontal="center"
            )


            for column_index in range(10):

                total = sum(

                    float(
                        row[column_index]
                    )

                    if row[column_index] is not None

                    else 0

                    for row in report_data

                )


                cell = worksheet.cell(
                    row=current_row,
                    column=column_index + 2
                )


                cell.value = total


                cell.font = Font(
                    bold=True
                )


                cell.fill = PatternFill(
                    "solid",
                    fgColor=yellow
                )


                cell.border = border


                cell.alignment = Alignment(
                    horizontal="center"
                )


                if column_index % 2 == 1:

                    cell.number_format = "0.000"


            current_row += 3


        # -------------------------------------------------
        # عرض الأعمدة
        # -------------------------------------------------

        worksheet.column_dimensions["A"].width = 20


        for column in range(2, 12):

            worksheet.column_dimensions[
                get_column_letter(column)
            ].width = 16


        # اتجاه Excel من اليمين لليسار

        worksheet.sheet_view.rightToLeft = True


        # إعدادات الطباعة

        worksheet.page_setup.orientation = "landscape"

        worksheet.page_setup.paperSize = (
            worksheet.PAPERSIZE_A4
        )

        worksheet.page_setup.fitToWidth = 1

        worksheet.page_setup.fitToHeight = 0

        worksheet.sheet_properties.pageSetUpPr.fitToPage = True


        # تجميد جزء من الجدول

        worksheet.freeze_panes = "B5"


        # -------------------------------------------------
        # إنشاء الملف في الذاكرة
        # -------------------------------------------------

        file_stream = io.BytesIO()

        workbook.save(
            file_stream
        )

        file_stream.seek(0)


        filename = (
            "تقرير_"
            + report_date
            + ".xlsx"
        )


        return send_file(

            file_stream,

            as_attachment=True,

            download_name=filename,

            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )

        )


    except Exception as error:

        print(
            "EXCEL ERROR:",
            error
        )


        return jsonify({

            "success": False,

            "message":
                str(error)

        }), 500


# =========================================================
# تشغيل البرنامج
# =========================================================

if __name__ == "__main__":

    init_database()

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=False
    )